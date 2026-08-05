"""De onde as linhas da planilha vêm.

Separa *ler a planilha* de *interpretar as colunas*: as fontes do
LiveClin e do WebDiet só precisam de uma lista de dicionários, venha ela
de um arquivo em disco ou de uma aba no Google Sheets.
"""

from __future__ import annotations

import csv
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from .base import ErroDeFonte

ESCOPOS_SHEETS = ["https://www.googleapis.com/auth/spreadsheets.readonly"]

_INSTRUCAO_GOOGLE = (
    "Para ler direto do Google Sheets instale as dependências opcionais:\n"
    "  pip install -r requirements-google.txt\n"
    "e use o mesmo credentials.json do Google Calendar."
)


class Planilha(ABC):
    @abstractmethod
    def linhas(self) -> list[dict[str, Any]]:
        """Linhas da planilha, com o cabeçalho como chave."""

    @property
    @abstractmethod
    def descricao(self) -> str:
        """Como a planilha aparece nas mensagens de erro."""


class PlanilhaArquivo(Planilha):
    def __init__(self, caminho: str | Path) -> None:
        self.caminho = Path(caminho).expanduser()

    @property
    def descricao(self) -> str:
        return str(self.caminho)

    def linhas(self) -> list[dict[str, Any]]:
        if not self.caminho.exists():
            raise ErroDeFonte(f"planilha não encontrada em {self.caminho}")
        if self.caminho.suffix.lower() in (".xlsx", ".xlsm"):
            return self._do_xlsx()
        return self._do_csv()

    def _do_csv(self) -> list[dict[str, Any]]:
        with self.caminho.open("r", encoding="utf-8-sig", newline="") as arquivo:
            amostra = arquivo.read(8192)
            arquivo.seek(0)
            try:
                dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
            except csv.Error:
                # Exportação brasileira costuma usar ponto e vírgula.
                dialeto = csv.excel
                dialeto.delimiter = ";" if amostra.count(";") > amostra.count(",") else ","
            return list(csv.DictReader(arquivo, dialect=dialeto))

    def _do_xlsx(self) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as erro:  # pragma: no cover - depende do ambiente
            raise ErroDeFonte(
                "Para ler .xlsx instale o openpyxl (pip install openpyxl) ou "
                "exporte a planilha como CSV."
            ) from erro

        planilha = load_workbook(self.caminho, read_only=True, data_only=True).active
        linhas = planilha.iter_rows(values_only=True)
        try:
            cabecalho = [str(c) if c is not None else "" for c in next(linhas)]
        except StopIteration:
            return []
        return [dict(zip(cabecalho, linha)) for linha in linhas]


class PlanilhaGoogleSheets(Planilha):
    """Aba de uma planilha do Google, lida por OAuth.

    Usa as mesmas credenciais do Google Calendar. O agente nunca recebe a
    senha da conta: a autorização acontece no navegador da sua máquina e
    o token fica em disco local, com permissão somente de leitura.
    """

    def __init__(
        self,
        spreadsheet_id: str,
        aba: str | None = None,
        credenciais: str | Path = "credentials.json",
        token: str | Path = "token_sheets.json",
    ) -> None:
        if not spreadsheet_id:
            raise ErroDeFonte("informe o spreadsheet_id da planilha do Google")
        self.spreadsheet_id = spreadsheet_id
        self.aba = aba
        self.credenciais = Path(credenciais).expanduser()
        self.token = Path(token).expanduser()

    @property
    def descricao(self) -> str:
        alvo = f"{self.spreadsheet_id}"
        return f"Google Sheets {alvo}" + (f" (aba {self.aba})" if self.aba else "")

    def _servico(self):
        try:
            from google.auth.transport.requests import Request
            from google.oauth2.credentials import Credentials
            from google_auth_oauthlib.flow import InstalledAppFlow
            from googleapiclient.discovery import build
        except ImportError as erro:
            raise ErroDeFonte(_INSTRUCAO_GOOGLE) from erro

        credencial = None
        if self.token.exists():
            credencial = Credentials.from_authorized_user_file(
                str(self.token), ESCOPOS_SHEETS
            )
        if credencial is None or not credencial.valid:
            if credencial and credencial.expired and credencial.refresh_token:
                credencial.refresh(Request())
            else:
                if not self.credenciais.exists():
                    raise ErroDeFonte(
                        f"não encontrei {self.credenciais}.\n{_INSTRUCAO_GOOGLE}"
                    )
                fluxo = InstalledAppFlow.from_client_secrets_file(
                    str(self.credenciais), ESCOPOS_SHEETS
                )
                credencial = fluxo.run_local_server(port=0)
            self.token.write_text(credencial.to_json(), encoding="utf-8")

        return build("sheets", "v4", credentials=credencial)

    def linhas(self) -> list[dict[str, Any]]:
        servico = self._servico()
        intervalo = self.aba if self.aba else None
        try:
            resposta = (
                servico.spreadsheets()
                .values()
                .get(
                    spreadsheetId=self.spreadsheet_id,
                    range=intervalo or "A:ZZ",
                    valueRenderOption="FORMATTED_VALUE",
                )
                .execute()
            )
        except Exception as erro:  # a biblioteca do Google levanta HttpError
            raise ErroDeFonte(
                f"não consegui ler {self.descricao}: {erro}. "
                "Confira o spreadsheet_id e se a conta autorizada enxerga a planilha."
            ) from erro

        valores = resposta.get("values", [])
        return montar_linhas(valores)


def montar_linhas(valores: list[list[Any]]) -> list[dict[str, Any]]:
    """Converte a matriz do Sheets em dicionários por cabeçalho.

    O Sheets corta células vazias no fim de cada linha, então as linhas
    curtas são completadas para alinhar com o cabeçalho.
    """
    if not valores:
        return []
    cabecalho = [str(c).strip() for c in valores[0]]
    linhas: list[dict[str, Any]] = []
    for bruta in valores[1:]:
        completada = list(bruta) + [""] * (len(cabecalho) - len(bruta))
        linhas.append(dict(zip(cabecalho, completada)))
    return linhas
