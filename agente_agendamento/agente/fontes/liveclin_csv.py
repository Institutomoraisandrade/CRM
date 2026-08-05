"""Leitura da planilha de pacientes exportada do LiveClin.

O LiveClin não oferece API pública, então o caminho suportado é a
exportação em CSV/XLSX. Os nomes de coluna variam entre exportações, por
isso cada campo aceita vários apelidos e pode ser sobrescrito na
configuração (seção ``[fonte.colunas]``).
"""

from __future__ import annotations

import csv
from datetime import date, timedelta
from pathlib import Path
from typing import Any, Iterable

from ..modelos import Paciente, StatusPaciente
from ..planos import CATALOGO_PADRAO, Plano, resolver_plano
from ..util import chave, ler_data, so_digitos
from .base import ErroDeFonte, FontePacientes

APELIDOS_COLUNA: dict[str, tuple[str, ...]] = {
    "nome": ("nome", "paciente", "nome do paciente", "nome completo"),
    "plano": ("plano", "plano selecionado", "tipo de plano", "plano contratado"),
    "plano_inicio": (
        "inicio do plano",
        "data de inicio",
        "data inicio",
        "inicio",
        "inicio do acompanhamento",
        "data de contratacao",
    ),
    "plano_fim": (
        "fim do plano",
        "data de fim",
        "data fim",
        "vencimento",
        "data de vencimento",
        "validade",
        "vigencia",
        "termino",
    ),
    "status": ("status", "situacao", "tipo de cliente"),
    "whatsapp": ("whatsapp", "telefone", "celular", "whatsapp do paciente"),
    "email": ("email", "e-mail"),
    "ultima_consulta": (
        "ultima consulta",
        "data da ultima consulta",
        "ultimo atendimento",
        "ultima avaliacao",
    ),
    "etiquetas": ("etiquetas", "tags", "marcadores"),
    "id_externo": ("id", "codigo", "id do paciente"),
}

APELIDOS_STATUS: dict[str, StatusPaciente] = {
    "ativo": StatusPaciente.ATIVO,
    "ativos": StatusPaciente.ATIVO,
    "em acompanhamento": StatusPaciente.ATIVO,
    "pausado": StatusPaciente.PAUSADO,
    "pausados": StatusPaciente.PAUSADO,
    "em pausa": StatusPaciente.PAUSADO,
    "inativo": StatusPaciente.INATIVO,
    "inativos": StatusPaciente.INATIVO,
    "finalizado": StatusPaciente.INATIVO,
    "encerrado": StatusPaciente.INATIVO,
    "cancelado": StatusPaciente.INATIVO,
}


def _ler_linhas(caminho: Path) -> list[dict[str, Any]]:
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        return _ler_xlsx(caminho)
    return _ler_csv(caminho)


def _ler_csv(caminho: Path) -> list[dict[str, Any]]:
    with caminho.open("r", encoding="utf-8-sig", newline="") as arquivo:
        amostra = arquivo.read(8192)
        arquivo.seek(0)
        try:
            dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
        except csv.Error:
            # Exportação brasileira costuma usar ponto e vírgula.
            dialeto = csv.excel
            dialeto.delimiter = ";" if amostra.count(";") > amostra.count(",") else ","
        return list(csv.DictReader(arquivo, dialect=dialeto))


def _ler_xlsx(caminho: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as erro:  # pragma: no cover - depende do ambiente
        raise ErroDeFonte(
            "Para ler .xlsx instale o openpyxl (pip install openpyxl) ou "
            "exporte a planilha do LiveClin como CSV."
        ) from erro

    planilha = load_workbook(caminho, read_only=True, data_only=True).active
    linhas = planilha.iter_rows(values_only=True)
    try:
        cabecalho = [str(c) if c is not None else "" for c in next(linhas)]
    except StopIteration:
        return []
    return [dict(zip(cabecalho, linha)) for linha in linhas]


class FonteLiveClinCSV(FontePacientes):
    def __init__(
        self,
        caminho: str | Path,
        colunas: dict[str, str] | None = None,
        catalogo: dict[str, Plano] | None = None,
    ) -> None:
        self.caminho = Path(caminho).expanduser()
        self.colunas_config = {campo: chave(col) for campo, col in (colunas or {}).items()}
        self.catalogo = catalogo if catalogo is not None else dict(CATALOGO_PADRAO)
        self._avisos: list[str] = []

    @property
    def avisos(self) -> list[str]:
        return self._avisos

    def _mapear_cabecalho(self, cabecalhos: Iterable[str]) -> dict[str, str]:
        disponiveis = {chave(c): c for c in cabecalhos if c}
        mapa: dict[str, str] = {}
        for campo, apelidos in APELIDOS_COLUNA.items():
            configurado = self.colunas_config.get(campo)
            if configurado and configurado in disponiveis:
                mapa[campo] = disponiveis[configurado]
                continue
            for apelido in apelidos:
                if apelido in disponiveis:
                    mapa[campo] = disponiveis[apelido]
                    break
        return mapa

    def carregar(self) -> list[Paciente]:
        if not self.caminho.exists():
            raise ErroDeFonte(
                f"planilha do LiveClin não encontrada em {self.caminho}. "
                "Exporte os pacientes no LiveClin e aponte 'fonte.caminho' para o arquivo."
            )

        linhas = _ler_linhas(self.caminho)
        if not linhas:
            raise ErroDeFonte(f"a planilha {self.caminho} está vazia.")

        mapa = self._mapear_cabecalho(linhas[0].keys())
        faltando = [c for c in ("nome", "plano") if c not in mapa]
        if faltando:
            raise ErroDeFonte(
                "não encontrei as colunas obrigatórias "
                + ", ".join(faltando)
                + f" na planilha {self.caminho}. Colunas lidas: "
                + ", ".join(str(c) for c in linhas[0].keys())
                + ". Configure os nomes reais em [fonte.colunas]."
            )

        pacientes: list[Paciente] = []
        for numero, linha in enumerate(linhas, start=2):
            try:
                paciente = self._montar_paciente(linha, mapa)
            except ValueError as erro:
                self._avisos.append(f"linha {numero} ignorada: {erro}")
                continue
            if paciente is not None:
                pacientes.append(paciente)
        return pacientes

    def _valor(self, linha: dict[str, Any], mapa: dict[str, str], campo: str) -> str | None:
        coluna = mapa.get(campo)
        if coluna is None:
            return None
        bruto = linha.get(coluna)
        if bruto is None:
            return None
        texto = str(bruto).strip()
        return texto or None

    def _montar_paciente(
        self, linha: dict[str, Any], mapa: dict[str, str]
    ) -> Paciente | None:
        nome = self._valor(linha, mapa, "nome")
        if not nome:
            return None

        plano_bruto = self._valor(linha, mapa, "plano")
        if not plano_bruto:
            raise ValueError(f"{nome} está sem plano preenchido")
        plano = resolver_plano(plano_bruto, self.catalogo)
        if plano is None:
            raise ValueError(
                f"plano {plano_bruto!r} de {nome} não está no catálogo "
                "(configure-o em [planos])"
            )

        inicio = ler_data(self._valor(linha, mapa, "plano_inicio"))
        fim = ler_data(self._valor(linha, mapa, "plano_fim"))
        if inicio is None and fim is None:
            raise ValueError(
                f"{nome} está sem data de início e sem data de fim do plano"
            )
        if inicio is None:
            inicio = fim - timedelta(days=plano.duracao_dias)
        if fim is None:
            fim = plano.data_fim(inicio)
        if fim < inicio:
            raise ValueError(
                f"{nome} tem fim do plano ({fim:%d/%m/%Y}) anterior ao início "
                f"({inicio:%d/%m/%Y})"
            )

        status_bruto = self._valor(linha, mapa, "status")
        status = StatusPaciente.ATIVO
        if status_bruto:
            status = APELIDOS_STATUS.get(chave(status_bruto), StatusPaciente.ATIVO)

        ultima = ler_data(self._valor(linha, mapa, "ultima_consulta"))
        if ultima is not None and ultima > date.today():
            raise ValueError(
                f"{nome} tem última consulta no futuro ({ultima:%d/%m/%Y})"
            )

        etiquetas_bruto = self._valor(linha, mapa, "etiquetas") or ""
        etiquetas = [e.strip() for e in etiquetas_bruto.replace(";", ",").split(",") if e.strip()]

        return Paciente(
            nome=nome,
            plano=plano,
            plano_inicio=inicio,
            plano_fim=fim,
            status=status,
            whatsapp=so_digitos(self._valor(linha, mapa, "whatsapp")),
            email=self._valor(linha, mapa, "email"),
            ultima_consulta=ultima,
            etiquetas=etiquetas,
            id_externo=self._valor(linha, mapa, "id_externo"),
        )
